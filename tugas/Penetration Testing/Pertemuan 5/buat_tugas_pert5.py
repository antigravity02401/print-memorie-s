"""
Script untuk membuat dokumen Tugas Pertemuan 5 Metodologi Penelitian
Nama: M. Arifin Ilham | NIM: 6404240008
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

def set_cell_bg(cell, hex_color):
    """Set background color of a table cell."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def set_cell_border(cell, **kwargs):
    """Set border for a table cell."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), '2B5BA8')
        tcBorders.append(border)
    tcPr.append(tcBorders)

def add_heading_kop(doc):
    """Add letterhead."""
    # Kop surat
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("KEMENTERIAN PENDIDIKAN, TINGGI, SAINS, DAN TEKNOLOGI")
    run.bold = True
    run.font.size = Pt(11)
    run.font.name = 'Times New Roman'

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = p2.add_run("POLITEKNIK NEGERI BENGKALIS")
    run2.bold = True
    run2.font.size = Pt(12)
    run2.font.name = 'Times New Roman'

    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run3 = p3.add_run("Jalan Bathin Alam, Sungai Alam, Bengkalis, Riau 28711")
    run3.font.size = Pt(10)
    run3.font.name = 'Times New Roman'
    
    p4 = doc.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run4 = p4.add_run("Telepon: (+62766) 24566  |  Fax: (+62766) 800 1000  |  www.polbeng.ac.id")
    run4.font.size = Pt(10)
    run4.font.name = 'Times New Roman'

    # Horizontal line
    p5 = doc.add_paragraph()
    p5.paragraph_format.space_before = Pt(4)
    p5.paragraph_format.space_after = Pt(4)
    border = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '12')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '1F3D7A')
    border.append(bottom)
    p5._p.get_or_add_pPr().append(border)

def add_title_block(doc):
    """Add title and identity block."""
    doc.add_paragraph()
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("TUGAS PERTEMUAN 5")
    run.bold = True
    run.font.size = Pt(14)
    run.font.name = 'Times New Roman'
    run.font.color.rgb = RGBColor(0x1F, 0x3D, 0x7A)
    
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = p2.add_run("Perbandingan Pendekatan Penelitian Terkait Topik Tugas Akhir")
    run2.bold = True
    run2.font.size = Pt(13)
    run2.font.name = 'Times New Roman'
    run2.font.color.rgb = RGBColor(0x1F, 0x3D, 0x7A)
    
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run3 = p3.add_run("Metodologi Penelitian (KSI252509) — Program Studi Sarjana Terapan Keamanan Sistem Informasi")
    run3.font.size = Pt(10)
    run3.font.name = 'Times New Roman'
    run3.font.italic = True
    
    doc.add_paragraph()

    # Identity table
    tbl = doc.add_table(rows=5, cols=2)
    tbl.style = 'Table Grid'
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

    identitas = [
        ("Nama Mahasiswa", "M. Arifin Ilham"),
        ("NIM", "6404240008"),
        ("Kelas", "5A"),
        ("Bidang Peminatan", "Keamanan Informasi"),
        ("Topik / Sub-Minat", "Pengujian Keamanan Aplikasi Web (Penetration Testing pada nyamankos.ours.web.id)"),
    ]

    for i, (key, val) in enumerate(identitas):
        row = tbl.rows[i]
        # Key cell
        kc = row.cells[0]
        kc.width = Cm(5)
        set_cell_bg(kc, "E8EDF7")
        kp = kc.paragraphs[0]
        kp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        kr = kp.add_run(key)
        kr.bold = True
        kr.font.size = Pt(10)
        kr.font.name = 'Times New Roman'
        
        # Value cell
        vc = row.cells[1]
        vp = vc.paragraphs[0]
        vr = vp.add_run(val)
        vr.font.size = Pt(10)
        vr.font.name = 'Times New Roman'

    doc.add_paragraph()

def add_section_heading(doc, text, number):
    """Add a formatted section heading."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(f"{number}. {text}")
    run.bold = True
    run.font.size = Pt(12)
    run.font.name = 'Times New Roman'
    run.font.color.rgb = RGBColor(0x1F, 0x3D, 0x7A)
    
    # Add underline-like bottom border
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '2B5BA8')
    pBdr.append(bottom)
    pPr.append(pBdr)

def build_comparison_table(doc):
    """Build the main comparison table."""
    # Columns: Aspek | Kuantitatif | Kualitatif | Campuran | DSR
    headers = ["Aspek", "Kuantitatif", "Kualitatif", "Campuran\n(Mixed Methods)", "Design Science\nResearch (DSR)"]
    
    rows_data = [
        (
            "Definisi\n(modul)",
            "Mengukur variabel dengan angka dan menguji hubungan/perbedaan antar variabel menggunakan analisis statistik.",
            "Memahami makna, pengalaman, dan konteks suatu fenomena secara mendalam melalui data naratif.",
            "Menggabungkan data kuantitatif dan kualitatif; angka menunjukkan \"apa\", narasi menjelaskan \"mengapa\".",
            "Merancang dan mengevaluasi artefak (sistem, model, metode, atau purwarupa) untuk memecahkan masalah praktis."
        ),
        (
            "Kata kerja\npetunjuk",
            "Mengukur, membandingkan, menguji pengaruh",
            "Memahami, menggali, mengapa, bagaimana pengalaman",
            "Mengukur sekaligus menggali alasan",
            "Merancang, membangun, mengembangkan artefak/sistem"
        ),
        (
            "Teknik\npengumpulan\ndata",
            "• Survei/kuesioner berskala\n• Eksperimen terkontrol\n• Log sistem (data numerik)",
            "• Wawancara mendalam/semi-terstruktur\n• Observasi partisipatif\n• Studi dokumen dan studi kasus",
            "• Gabungan kuantitatif & kualitatif\n• Sekuensial (bertahap) atau Konkuren (bersamaan)",
            "• Identifikasi masalah\n• Rancang artefak\n• Evaluasi artefak (iteratif)"
        ),
        (
            "Teknik\nanalisis",
            "Statistik deskriptif & inferensial (SPSS, R, Python): uji-t, regresi, ANOVA",
            "Analisis tematik, triangulasi sumber (NVivo, Atlas.ti, manual)",
            "Triangulasi data angka dan narasi; saling memverifikasi",
            "Evaluasi artefak terhadap kriteria yang ditetapkan (akurasi, presisi, efektivitas)"
        ),
        (
            "Bentuk data /\nluaran pada\ntopik ini",
            "Jumlah temuan per kategori WSTG, skor CVSS, tabel risiko kuantitatif",
            "Narasi hasil uji, dampak kerentanan, analisis konteks aplikasi kos",
            "Skor CVSS + survei/wawancara pengembang tentang kendala perbaikan",
            "Artefak: dokumen rekomendasi mitigasi yang divalidasi melalui uji ulang"
        ),
        (
            "Contoh di\nmodul",
            "MFA menurunkan insiden akses tidak sah (diuji dengan uji-t berpasangan)",
            "Mengapa pegawai tetap klik phishing meski sudah pelatihan",
            "Survei kepatuhan kebijakan keamanan lalu wawancara pegawai patuh-rendah",
            "Sistem deteksi phishing berbasis ML, diukur akurasi, presisi, recall"
        ),
    ]

    # Create table with 5 columns
    tbl = doc.add_table(rows=len(rows_data) + 1, cols=5)
    tbl.style = 'Table Grid'
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Column widths
    col_widths = [Cm(3.2), Cm(3.8), Cm(3.8), Cm(3.8), Cm(3.8)]
    for row in tbl.rows:
        for i, cell in enumerate(row.cells):
            cell.width = col_widths[i]

    # Header row
    header_row = tbl.rows[0]
    header_colors = ["1F3D7A", "1F3D7A", "1F3D7A", "1F3D7A", "1F3D7A"]
    for i, (cell, header, color) in enumerate(zip(header_row.cells, headers, header_colors)):
        set_cell_bg(cell, color)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(header)
        run.bold = True
        run.font.size = Pt(9)
        run.font.name = 'Times New Roman'
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # Data rows
    row_bg_colors = ["F0F4FF", "FFFFFF", "F0F4FF", "FFFFFF", "F0F4FF", "FFFFFF"]
    for row_idx, (row_data, bg) in enumerate(zip(rows_data, row_bg_colors)):
        tbl_row = tbl.rows[row_idx + 1]
        for col_idx, (cell, text) in enumerate(zip(tbl_row.cells, row_data)):
            set_cell_bg(cell, bg if col_idx > 0 else "D0DAF0")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(text)
            run.font.size = Pt(8.5)
            run.font.name = 'Times New Roman'
            if col_idx == 0:
                run.bold = True
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    doc.add_paragraph()

def add_kelebihan_keterbatasan(doc):
    """Add section 2: Kelebihan dan Keterbatasan."""
    
    # Rumusan masalah
    p = doc.add_paragraph()
    run = p.add_run("Rumusan Masalah Topik: ")
    run.bold = True
    run.font.size = Pt(10)
    run.font.name = 'Times New Roman'
    run2 = p.add_run(
        "\"Sejauh mana aplikasi nyamankos (https://nyamankos.ours.web.id) rentan terhadap kerentanan "
        "kategori OWASP Top 10, dan bagaimana mitigasi yang sesuai untuk meningkatkan keamanannya?\""
    )
    run2.font.size = Pt(10)
    run2.font.name = 'Times New Roman'
    run2.italic = True
    p.paragraph_format.space_after = Pt(8)

    pendekatan_data = [
        {
            "judul": "Kuantitatif",
            "warna_judul": "1565C0",
            "kelebihan": [
                "Bagian \"sejauh mana rentan\" terjawab terukur melalui jumlah temuan per kategori WSTG dan skor CVSS.",
                "Hasilnya objektif, dapat dikuantifikasi, dan mudah diprioritaskan oleh tim pengembang.",
                "Data log hasil uji (HTTP response, status code, payload berhasil/gagal) tersedia secara alami dalam proses pentest."
            ],
            "keterbatasan": [
                "Bagian \"bagaimana mitigasi yang sesuai\" tidak terjawab oleh angka semata — pendekatan ini tidak menghasilkan solusi.",
                "Desain eksperimen (perlakuan & pembanding) tidak ada; sampelnya hanya satu aplikasi sehingga generalisasi sulit.",
                "Tidak menjelaskan konteks atau alasan mengapa kerentanan muncul dalam sistem ini."
            ]
        },
        {
            "judul": "Kualitatif",
            "warna_judul": "2E7D32",
            "kelebihan": [
                "Cocok menjelaskan akar masalah kerentanan, skenario serangan, dan dampak ke konteks aplikasi kos (data penghuni, data booking).",
                "Studi kasus pada satu aplikasi nyata sesuai dengan objek penelitian yang spesifik.",
                "Dapat menggali perspektif tim pengembang mengenai penyebab kerentanan muncul."
            ],
            "keterbatasan": [
                "\"Sejauh mana rentan\" menuntut ukuran risiko objektif — narasi saja rentan bias interpretasi.",
                "Tidak menghasilkan solusi mitigasi yang tervalidasi; hasil hanya deskriptif.",
                "Sumber data naratif terbatas; hanya ada tim pengembang yang dapat diwawancarai."
            ]
        },
        {
            "judul": "Campuran (Mixed Methods)",
            "warna_judul": "6A1FAA",
            "kelebihan": [
                "Dapat menjawab \"apa\" (skor risiko kuantitatif) dan \"mengapa\" (konteks kualitatif) secara bersamaan.",
                "Triangulasi data angka dan narasi meningkatkan validitas temuan.",
                "Fleksibel: data kuantitatif dari pentest, data kualitatif dari wawancara pengembang."
            ],
            "keterbatasan": [
                "Sumber data naratif sangat terbatas — hanya tim pengembang; kurang representatif untuk pendekatan campuran yang kuat.",
                "Butuh waktu dan sumber daya lebih besar; jadwal Minggu 5–15 sudah sangat padat.",
                "Berisiko jatuh ke kesalahan umum 'memaksakan campuran' (slide 22 modul) tanpa alasan metodologis yang kuat."
            ]
        },
        {
            "judul": "Design Science Research (DSR)",
            "warna_judul": "B71C1C",
            "kelebihan": [
                "Siklus DSR (identifikasi masalah → rancang artefak → evaluasi) terpetakan langsung pada alur pentest: temuan kerentanan = masalah, rekomendasi mitigasi = artefak, uji ulang = evaluasi.",
                "Menghasilkan solusi nyata yang dapat langsung dipakai tim pengembang sebelum diserahkan ke mitra.",
                "Contoh latihan slide 21 modul secara eksplisit memetakan rumusan masalah 'SQL Injection dan mitigasinya' ke DSR — identik dengan topik ini."
            ],
            "keterbatasan": [
                "Evaluasi artefak bisa subjektif bila kriteria efektivitas tidak ditetapkan sejak awal (peringatan slide 22 modul).",
                "Uji ulang bergantung pada kesediaan tim pengembang menerapkan perbaikan dalam sisa waktu perkuliahan.",
                "Siklus iteratif memerlukan komitmen waktu tambahan di luar jadwal pentest yang sudah padat."
            ]
        }
    ]

    colors_bg = {
        "Kuantitatif": ("E3F2FD", "1565C0"),
        "Kualitatif": ("E8F5E9", "2E7D32"),
        "Campuran (Mixed Methods)": ("F3E5F5", "6A1FAA"),
        "Design Science Research (DSR)": ("FFEBEE", "B71C1C"),
    }

    for pd in pendekatan_data:
        judul = pd["judul"]
        bg, title_color = colors_bg.get(judul, ("F5F5F5", "333333"))
        
        # Create box-like paragraph for heading
        p_head = doc.add_paragraph()
        p_head.paragraph_format.space_before = Pt(8)
        p_head.paragraph_format.space_after = Pt(2)
        p_head.paragraph_format.left_indent = Cm(0.3)
        r = p_head.add_run(f"◆  {judul}")
        r.bold = True
        r.font.size = Pt(11)
        r.font.name = 'Times New Roman'
        r_color = tuple(int(title_color[i:i+2], 16) for i in (0, 2, 4))
        r.font.color.rgb = RGBColor(*r_color)

        # Kelebihan
        pk = doc.add_paragraph()
        pk.paragraph_format.space_before = Pt(2)
        pk.paragraph_format.space_after = Pt(1)
        pk.paragraph_format.left_indent = Cm(0.8)
        rk = pk.add_run("Kelebihan:")
        rk.bold = True
        rk.font.size = Pt(10)
        rk.font.name = 'Times New Roman'
        rk.font.color.rgb = RGBColor(0x2E, 0x7D, 0x32)

        for point in pd["kelebihan"]:
            pb = doc.add_paragraph(style='List Bullet')
            pb.paragraph_format.left_indent = Cm(1.5)
            pb.paragraph_format.space_before = Pt(0)
            pb.paragraph_format.space_after = Pt(1)
            rb = pb.add_run(point)
            rb.font.size = Pt(10)
            rb.font.name = 'Times New Roman'

        # Keterbatasan
        pkt = doc.add_paragraph()
        pkt.paragraph_format.space_before = Pt(3)
        pkt.paragraph_format.space_after = Pt(1)
        pkt.paragraph_format.left_indent = Cm(0.8)
        rkt = pkt.add_run("Keterbatasan:")
        rkt.bold = True
        rkt.font.size = Pt(10)
        rkt.font.name = 'Times New Roman'
        rkt.font.color.rgb = RGBColor(0xC6, 0x28, 0x28)

        for point in pd["keterbatasan"]:
            pb2 = doc.add_paragraph(style='List Bullet')
            pb2.paragraph_format.left_indent = Cm(1.5)
            pb2.paragraph_format.space_before = Pt(0)
            pb2.paragraph_format.space_after = Pt(1)
            rb2 = pb2.add_run(point)
            rb2.font.size = Pt(10)
            rb2.font.name = 'Times New Roman'

    doc.add_paragraph()

def add_pilihan_pendekatan(doc):
    """Add section 3: Chosen Approach."""
    
    # Box/callout style
    tbl = doc.add_table(rows=1, cols=1)
    tbl.style = 'Table Grid'
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.rows[0].cells[0]
    set_cell_bg(cell, "FFF8E1")
    
    p0 = cell.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r0 = p0.add_run("✔  Pendekatan Terpilih: Design Science Research (DSR)")
    r0.bold = True
    r0.font.size = Pt(12)
    r0.font.name = 'Times New Roman'
    r0.font.color.rgb = RGBColor(0xB7, 0x1C, 0x1C)

    doc.add_paragraph()

    alasan_paras = [
        (
            "Kesesuaian Kata Kerja pada Rumusan Masalah",
            "Rumusan masalah mengandung frasa \"bagaimana mitigasi yang sesuai\" dan tujuan penelitian mencakup "
            "\"merancang rekomendasi mitigasi\". Menurut panduan praktis modul (slide 20), kata kerja merancang "
            "dan mengembangkan mengarahkan langsung ke DSR."
        ),
        (
            "Pemetaan Langsung ke Contoh Latihan Modul",
            "Latihan slide 21 modul secara eksplisit memetakan rumusan masalah tentang kerentanan SQL Injection "
            "dan mitigasinya ke DSR. Topik penelitian ini identik secara struktur: mengidentifikasi kerentanan "
            "OWASP Top 10 kemudian merumuskan mitigasinya."
        ),
        (
            "Siklus DSR Terpetakan pada Rencana Kerja yang Ada",
            "Fase identifikasi masalah terpenuhi melalui Intelligence Gathering dan Vulnerability Analysis (Mg 5–9). "
            "Fase rancang artefak terpenuhi melalui penyusunan rekomendasi mitigasi (Mg 14–15). "
            "Fase evaluasi artefak terpenuhi melalui uji ulang atau verifikasi ke tim pengembang."
        ),
        (
            "Kriteria Evaluasi yang Terukur (Menghindari Subjektivitas)",
            "Sesuai peringatan modul (slide 22), kriteria evaluasi ditetapkan sejak awal: "
            "(a) penurunan skor CVSS rata-rata setelah perbaikan diterapkan, atau "
            "(b) jumlah temuan yang berstatus 'mitigasi terverifikasi' vs total temuan awal. "
            "Hal ini mencegah klaim 'efektif' tanpa bukti."
        ),
        (
            "Pengukuran Kuantitatif sebagai Data Pendukung, Bukan Pendekatan Utama",
            "Data kuantitatif (skor CVSS, jumlah temuan per kategori WSTG) tetap dipakai, "
            "tetapi berperan sebagai data input tahap identifikasi masalah dan metrik evaluasi artefak — "
            "bukan sebagai pendekatan utama. Ini menjaga konsistensi dengan pilihan DSR."
        ),
    ]

    for i, (sub_judul, isi) in enumerate(alasan_paras, 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(f"{i}. {sub_judul}")
        r.bold = True
        r.font.size = Pt(10.5)
        r.font.name = 'Times New Roman'
        r.font.color.rgb = RGBColor(0x1F, 0x3D, 0x7A)

        p2 = doc.add_paragraph()
        p2.paragraph_format.left_indent = Cm(0.8)
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(4)
        r2 = p2.add_run(isi)
        r2.font.size = Pt(10)
        r2.font.name = 'Times New Roman'

    doc.add_paragraph()

def add_referensi(doc):
    """Add references section."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    run = p.add_run("Referensi")
    run.bold = True
    run.font.size = Pt(11)
    run.font.name = 'Times New Roman'
    run.font.color.rgb = RGBColor(0x1F, 0x3D, 0x7A)

    refs = [
        "Creswell, J. W., & Creswell, J. D. (2018). Research Design: Qualitative, Quantitative, and Mixed Methods Approaches (5th ed.). Sage Publications.",
        "Wieringa, R. J. (2014). Design Science Methodology for Information Systems and Software Engineering. Springer.",
        "OWASP Foundation. (2020). OWASP Web Security Testing Guide (WSTG) v4.2. https://owasp.org/www-project-web-security-testing-guide/",
        "Hidayasari, N. (2026). Modul Metodologi Penelitian — Pertemuan 5: Jenis dan Pendekatan Penelitian. Politeknik Negeri Bengkalis.",
    ]

    for i, ref in enumerate(refs, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.first_line_indent = Cm(-0.8)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(f"[{i}]  {ref}")
        r.font.size = Pt(9.5)
        r.font.name = 'Times New Roman'


def main():
    doc = Document()
    
    # Page margins
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(3)
        section.right_margin = Cm(2)

    # Set default font
    doc.styles['Normal'].font.name = 'Times New Roman'
    doc.styles['Normal'].font.size = Pt(11)

    # Build document
    add_heading_kop(doc)
    add_title_block(doc)

    add_section_heading(doc, "Tabel Perbandingan Keempat Pendekatan Penelitian", "1")
    p_note = doc.add_paragraph()
    rn = p_note.add_run(
        "Tabel berikut membandingkan empat pendekatan penelitian (Kuantitatif, Kualitatif, Campuran, dan "
        "Design Science Research) dikaitkan langsung dengan topik: Pengujian Keamanan Aplikasi Web "
        "nyamankos menggunakan OWASP Top 10 dan PTES."
    )
    rn.font.size = Pt(10)
    rn.font.name = 'Times New Roman'
    rn.italic = True
    p_note.paragraph_format.space_after = Pt(6)
    
    build_comparison_table(doc)

    add_section_heading(doc, "Kelebihan dan Keterbatasan Setiap Pendekatan pada Rumusan Masalah", "2")
    add_kelebihan_keterbatasan(doc)

    add_section_heading(doc, "Pendekatan Awal yang Dipilih dan Alasan", "3")
    add_pilihan_pendekatan(doc)

    add_referensi(doc)

    # Footer
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    p_footer = doc.add_paragraph()
    p_footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_footer.paragraph_format.space_before = Pt(12)
    
    # Border top
    pPr = p_footer._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    top = OxmlElement('w:top')
    top.set(qn('w:val'), 'single')
    top.set(qn('w:sz'), '6')
    top.set(qn('w:space'), '1')
    top.set(qn('w:color'), '1F3D7A')
    pBdr.append(top)
    pPr.append(pBdr)
    
    r_f = p_footer.add_run("© Tugas Pertemuan 5 Metodologi Penelitian  |  M. Arifin Ilham (6404240008)  |  Kelas 5A  |  Politeknik Negeri Bengkalis  |  2026")
    r_f.font.size = Pt(8)
    r_f.font.name = 'Times New Roman'
    r_f.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    # Save
    output_path = r"Tugas_Pert5_Metopen_M_Arifin_Ilham.docx"
    doc.save(output_path)
    print(f"Dokumen berhasil disimpan: {output_path}")

if __name__ == "__main__":
    main()
